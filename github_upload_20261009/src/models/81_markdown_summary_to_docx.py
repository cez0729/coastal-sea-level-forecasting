from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "增强数据后实验结果总结与下一步说明.md"
OUTPUT = ROOT / "增强数据后实验结果总结与下一步说明.docx"


def set_run_font(run, name: str = "Microsoft YaHei", size: int | None = None, bold: bool | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa: int = 9360):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    table.autofit = False


def add_formatted_paragraph(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Normal"]
    add_inline_runs(paragraph, text)
    return paragraph


def add_inline_runs(paragraph, text: str):
    # Split on inline code and bold markers; enough for this report's Markdown.
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Consolas", 9)
            run.font.color.rgb = RGBColor(80, 80, 80)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, bold=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    table_lines = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
        table_lines.append(lines[i].strip())
        i += 1
    rows = []
    for idx, line in enumerate(table_lines):
        parts = [part.strip() for part in line.strip("|").split("|")]
        if idx == 1 and all(set(part.replace(":", "")) <= {"-"} for part in parts):
            continue
        rows.append(parts)
    return rows, i


def add_markdown_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_width(table)

    # Allocate wider columns to narrative-heavy tables, compact widths to numeric columns.
    if cols == 2:
        widths = [2600, 6760]
    elif cols == 3:
        widths = [2100, 2600, 4660]
    elif cols == 4:
        widths = [1200, 3000, 2580, 2580]
    elif cols == 5:
        widths = [1050, 3200, 1700, 1700, 1710]
    elif cols == 6:
        widths = [850, 2400, 1100, 1350, 1350, 2310]
    else:
        base = 9360 // cols
        widths = [base] * cols

    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths[min(c_idx, len(widths) - 1)])
            text = row[c_idx] if c_idx < len(row) else ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 or re.fullmatch(r"[-+0-9. ±%]+", text) else WD_ALIGN_PARAGRAPH.LEFT
            add_inline_runs(p, text)
            for run in p.runs:
                set_run_font(run, size=9, bold=(r_idx == 0))
            if r_idx == 0:
                set_cell_shading(cell, "E8EEF5")
    doc.add_paragraph()


def add_code_block(doc: Document, block: list[str]):
    text = "\n".join(block)
    paragraph = doc.add_paragraph()
    paragraph.style = "CodeBlock"
    run = paragraph.add_run(text)
    set_run_font(run, "Consolas", 9)


def build_docx():
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    code_style = styles.add_style("CodeBlock", 1)
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code_style.font.size = Pt(9)
    code_style.paragraph_format.left_indent = Inches(0.18)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(8)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("海平面 residual forecasting 实验总结")
    set_run_font(r, size=9)
    r.font.color.rgb = RGBColor(90, 90, 90)

    first_title = True
    i = 0
    in_code = False
    code_lines: list[str] = []
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not stripped:
            i += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            rows, next_i = parse_table(lines, i)
            add_markdown_table(doc, rows)
            i = next_i
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            title = stripped[level:].strip()
            if first_title:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(title)
                set_run_font(run, size=20, bold=True)
                run.font.color.rgb = RGBColor.from_string("0B2545")
                doc.add_paragraph()
                first_title = False
            else:
                p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
                add_inline_runs(p, title)
                for run in p.runs:
                    set_run_font(run, size=16 if level == 1 else 13 if level == 2 else 12, bold=True)
            i += 1
            continue

        bullet_match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if bullet_match:
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, bullet_match.group(2))
            i += 1
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, stripped[2:])
            i += 1
            continue

        add_formatted_paragraph(doc, stripped)
        i += 1

    # Keep a blank paragraph at the end so Word has room after final code/table.
    doc.add_paragraph()
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    out = build_docx()
    print(out)
